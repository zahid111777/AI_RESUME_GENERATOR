"""
High-quality PDF  (ReportLab Platypus) and DOCX (python-docx) generators
that parse the structured resume HTML produced by the AI backend.
"""

import io
import os
import re
from bs4 import BeautifulSoup

# ── ReportLab ─────────────────────────────────────────────────────────────────
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_RIGHT
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer,
    HRFlowable, Table, TableStyle, KeepTogether,
)

# ── python-docx ───────────────────────────────────────────────────────────────
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── Colour palette ────────────────────────────────────────────────────────────
_NAVY    = HexColor('#1a3c6e')
_BLACK   = HexColor('#111111')
_DARK    = HexColor('#2a2a2a')
_GRAY    = HexColor('#555555')
_IGRAY   = HexColor('#3d3d3d')

_NAVY_RGB  = RGBColor(0x1a, 0x3c, 0x6e)
_BLACK_RGB = RGBColor(0x11, 0x11, 0x11)
_DARK_RGB  = RGBColor(0x2a, 0x2a, 0x2a)
_GRAY_RGB  = RGBColor(0x55, 0x55, 0x55)
_IGRAY_RGB = RGBColor(0x3d, 0x3d, 0x3d)

# ══════════════════════════════════════════════════════════════════════════════
#  Windows font registration for ReportLab
#  Falls back to Helvetica / Times-Roman if the TTF is not found.
# ══════════════════════════════════════════════════════════════════════════════
_WF = "C:/Windows/Fonts"
_REGISTERED: dict = {}


def _reg(name: str, reg: str, bld: str, ita: str, bi: str) -> bool:
    """Register a TrueType font family in ReportLab; return True on success."""
    try:
        for n, f in [(name, reg), (f"{name}-Bold", bld),
                     (f"{name}-Italic", ita), (f"{name}-BoldItalic", bi)]:
            pdfmetrics.registerFont(TTFont(n, os.path.join(_WF, f)))
        from reportlab.lib.fonts import addMapping
        addMapping(name, 0, 0, name)
        addMapping(name, 1, 0, f"{name}-Bold")
        addMapping(name, 0, 1, f"{name}-Italic")
        addMapping(name, 1, 1, f"{name}-BoldItalic")
        return True
    except Exception:
        return False


_REGISTERED['Calibri'] = _reg('Calibri', 'calibri.ttf',  'calibrib.ttf',
                               'calibrii.ttf', 'calibriz.ttf')
_REGISTERED['Georgia'] = _reg('Georgia', 'georgia.ttf',  'georgiab.ttf',
                               'georgiai.ttf', 'georgiaz.ttf')
_REGISTERED['Arial']   = _reg('Arial',   'arial.ttf',    'arialbd.ttf',
                               'ariali.ttf',   'arialbi.ttf')
_REGISTERED['Candara'] = _reg('Candara', 'Candara.ttf',  'Candarab.ttf',
                               'Candarai.ttf', 'Candaraz.ttf')


def _first_font(family_css: str) -> str:
    """Extract the first font name from a CSS font-family string."""
    # e.g. "'Montserrat', 'Candara', sans-serif" → "montserrat"
    m = re.match(r"[\'\"]?([\w\- ]+)[\'\"]?", family_css.strip())
    return m.group(1).strip().lower() if m else family_css.lower()


def _rl_fonts(family_css: str) -> tuple:
    """Return (regular, bold, italic, bold-italic) ReportLab font names."""
    first = _first_font(family_css)
    mapping = [
        ('montserrat', 'Arial'),   # closest Windows substitute for Montserrat
        ('calibri',    'Calibri'),
        ('candara',    'Candara'),
        ('georgia',    'Georgia'),
        ('arial',      'Arial'),
    ]
    for key, canon in mapping:
        if key in first and _REGISTERED.get(canon):
            return (canon, f"{canon}-Bold",
                    f"{canon}-Italic", f"{canon}-BoldItalic")
    # Built-in fallback
    if 'georgia' in first or 'serif' in first:
        return ('Times-Roman', 'Times-Bold',
                'Times-Italic', 'Times-BoldItalic')
    return ('Helvetica', 'Helvetica-Bold',
            'Helvetica-Oblique', 'Helvetica-BoldOblique')


def _docx_font(family_css: str) -> str:
    first = _first_font(family_css)
    if 'montserrat' in first: return 'Arial'
    if 'candara'    in first: return 'Candara'
    if 'georgia'    in first: return 'Georgia'
    if 'calibri'    in first: return 'Calibri'
    if 'arial'      in first: return 'Arial'
    return 'Calibri'


# ── Text helpers ──────────────────────────────────────────────────────────────
def _clean(text: str) -> str:
    return re.sub(r'\s+', ' ', text or '').strip()


def _get_text(el) -> str:
    return _clean(el.get_text()) if el else ''


# ══════════════════════════════════════════════════════════════════════════════
#  PDF  via  ReportLab Platypus
# ══════════════════════════════════════════════════════════════════════════════
def build_pdf(resume_html: str,
              heading_font_css: str = 'Calibri',
              body_font_css:    str = 'Calibri') -> bytes:

    soup = BeautifulSoup(resume_html, 'html.parser')
    buf  = io.BytesIO()

    _,   HFB, _,   _    = _rl_fonts(heading_font_css)
    BF,  BFB, BFI, _    = _rl_fonts(body_font_css)

    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=18*mm, rightMargin=18*mm,
        topMargin=15*mm,  bottomMargin=15*mm,
    )
    W = A4[0] - 36*mm   # usable text width

    # ── Styles ───────────────────────────────────────────────────────────────
    def S(name, **kw) -> ParagraphStyle:
        return ParagraphStyle(name, **kw)

    sName = S('Name', fontName=HFB, fontSize=20, leading=24,
              textColor=_NAVY, alignment=TA_CENTER,
              spaceBefore=0, spaceAfter=4, letterSpacing=1.5)

    sContact = S('Contact', fontName=BF, fontSize=9, leading=13,
                 textColor=_GRAY, alignment=TA_CENTER,
                 spaceBefore=0, spaceAfter=10)

    sSectTitle = S('SectTitle', fontName=BFB, fontSize=8.5, leading=11,
                   textColor=_NAVY, spaceBefore=10, spaceAfter=2,
                   letterSpacing=1.5)

    sJobTitle = S('JobTitle', fontName=BFB, fontSize=10.5, leading=13,
                  textColor=_BLACK, spaceBefore=0, spaceAfter=0)

    sDate = S('Date', fontName=BF, fontSize=9.5, leading=13,
              textColor=_GRAY, alignment=TA_RIGHT,
              spaceBefore=0, spaceAfter=0)

    sCompany = S('Company', fontName=BFI, fontSize=10, leading=12,
                 textColor=_IGRAY, spaceBefore=0, spaceAfter=2)

    sBullet = S('Bullet', fontName=BF, fontSize=10.5, leading=13,
                textColor=_DARK, leftIndent=10,
                spaceBefore=0, spaceAfter=1)

    sSkills = S('Skills', fontName=BF, fontSize=10.5, leading=14,
                textColor=_DARK, spaceBefore=0, spaceAfter=2)

    sSummary = S('Summary', fontName=BF, fontSize=10.5, leading=14,
                 textColor=_DARK, alignment=TA_JUSTIFY, spaceAfter=2)

    sSchool = S('School', fontName=BFI, fontSize=10, leading=12,
                textColor=_IGRAY, spaceBefore=0, spaceAfter=0)

    sGPA = S('GPA', fontName=BFB, fontSize=9.5, leading=12,
             textColor=_NAVY, alignment=TA_RIGHT,
             spaceBefore=0, spaceAfter=0)

    sCert = S('Cert', fontName=BF, fontSize=10.5, leading=13,
              textColor=_DARK, leftIndent=10,
              spaceBefore=0, spaceAfter=1)

    sProjName = S('ProjName', fontName=BFB, fontSize=10.5, leading=13,
                  textColor=_BLACK, spaceBefore=0, spaceAfter=0)

    sProjDesc = S('ProjDesc', fontName=BF, fontSize=10.5, leading=13,
                  textColor=_DARK, spaceBefore=0, spaceAfter=2)

    # ── Helpers ──────────────────────────────────────────────────────────────
    _two_ts = TableStyle([
        ('ALIGN',         (1, 0), (1, -1), 'RIGHT'),
        ('VALIGN',        (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING',   (0, 0), (-1, -1), 0),
        ('RIGHTPADDING',  (0, 0), (-1, -1), 0),
        ('TOPPADDING',    (0, 0), (-1, -1), 0),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
    ])

    def _two_col(left_p, right_p, lf=0.72):
        t = Table([[left_p, right_p]],
                  colWidths=[W * lf, W * (1 - lf)])
        t.setStyle(_two_ts)
        return t

    def _sect_hr():
        """Thin line directly under a section title (matches HTML border-bottom)."""
        return HRFlowable(width='100%', thickness=1.2,
                          color=_NAVY, spaceBefore=1, spaceAfter=5)

    def _header_hr():
        """Line after contact bar — sits cleanly below contact text, clear gap before sections."""
        return HRFlowable(width='100%', thickness=1.2,
                          color=_NAVY, spaceBefore=2, spaceAfter=12)

    # ── Story ─────────────────────────────────────────────────────────────────
    story = []

    hdr = soup.find(class_='resume-header')
    if hdr:
        h1 = hdr.find('h1')
        if h1:
            story.append(Paragraph(_get_text(h1).upper(), sName))
        cbar = hdr.find(class_='contact-bar')
        if cbar:
            parts = []
            for a in cbar.find_all('a'):
                txt  = _get_text(a)
                href = a.get('href', '')
                if txt:
                    parts.append(
                        f'<link href="{href}" color="#1a3c6e">{txt}</link>')
            if parts:
                story.append(Paragraph('  \u2022  '.join(parts), sContact))

    for sec in soup.find_all(class_='section'):
        title_el = sec.find(class_='section-title')
        if title_el:
            story.append(Paragraph(_get_text(title_el).upper(), sSectTitle))
            story.append(_sect_hr())

        sk = sec.find(class_='skills-text')
        if sk:
            story.append(Paragraph(_get_text(sk), sSkills))

        for p_el in sec.find_all('p', recursive=False):
            story.append(Paragraph(_get_text(p_el), sSummary))

        for job in sec.find_all(class_='job'):
            items = []
            jh = job.find(class_='job-header')
            if jh:
                items.append(_two_col(
                    Paragraph(_get_text(jh.find(class_='job-title')), sJobTitle),
                    Paragraph(_get_text(jh.find(class_='job-date')),  sDate),
                ))
            cmp = job.find(class_='job-company')
            if cmp:
                items.append(Paragraph(_get_text(cmp), sCompany))
            for li in job.find_all('li'):
                items.append(Paragraph(f'\u2022\u00a0\u00a0{_get_text(li)}', sBullet))
            items.append(Spacer(1, 6))
            story.append(KeepTogether(items))

        for edu in sec.find_all(class_='edu-entry'):
            items = []
            eh = edu.find(class_='edu-header')
            if eh:
                items.append(_two_col(
                    Paragraph(_get_text(eh.find(class_='edu-degree')), sJobTitle),
                    Paragraph(_get_text(eh.find(class_='edu-date')),   sDate),
                ))
            es = edu.find(class_='edu-subheader')
            if es:
                items.append(_two_col(
                    Paragraph(_get_text(es.find(class_='edu-school')), sSchool),
                    Paragraph(_get_text(es.find(class_='edu-gpa')),    sGPA),
                ))
            items.append(Spacer(1, 5))
            story.append(KeepTogether(items))

        for proj in sec.find_all(class_='proj-entry'):
            items = []
            ph = proj.find(class_='proj-header')
            if ph:
                ng    = ph.find(class_='proj-name-group') or ph
                pn    = _get_text(ng.find(class_='proj-name'))
                gh_el = ph.find(class_='proj-gh-link')
                pd    = _get_text(ph.find(class_='proj-date'))
                name_txt = pn
                if gh_el:
                    href   = gh_el.get('href', '')
                    gh_txt = _get_text(gh_el) or 'GitHub'
                    name_txt = (f'{pn}  <link href="{href}" color="#1a3c6e">'
                                f'<font size="8">{gh_txt}</font></link>')
                items.append(_two_col(
                    Paragraph(name_txt, sProjName),
                    Paragraph(pd, sDate),
                    lf=0.74,
                ))
            desc_el = proj.find(class_='proj-desc')
            if desc_el:
                items.append(Paragraph(_get_text(desc_el), sProjDesc))
            items.append(Spacer(1, 5))
            story.append(KeepTogether(items))

        for cert in sec.find_all(class_='cert-entry'):
            story.append(Paragraph(
                f'\u2022\u00a0\u00a0{_get_text(cert)}', sCert))
        story.append(Spacer(1, 3))

    doc.build(story)
    return buf.getvalue()


# ══════════════════════════════════════════════════════════════════════════════
#  DOCX  via  python-docx
#  All title↔date rows use a proper 2-column borderless Table —
#  no tab-stop tricks, dates are always right-aligned.
# ══════════════════════════════════════════════════════════════════════════════
def _d_spacing(p, before=0, after=0):
    pPr = p._p.get_or_add_pPr()
    sp  = OxmlElement('w:spacing')
    sp.set(qn('w:before'), str(before))
    sp.set(qn('w:after'),  str(after))
    pPr.append(sp)


def _d_border_bottom(p, color='1a3c6e', sz='10'):
    pPr  = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bot  = OxmlElement('w:bottom')
    bot.set(qn('w:val'),   'single')
    bot.set(qn('w:sz'),    sz)
    bot.set(qn('w:space'), '1')
    bot.set(qn('w:color'), color)
    pBdr.append(bot)
    pPr.append(pBdr)


def _d_run(p, text, font, size_pt,
           bold=False, italic=False, color=_DARK_RGB):
    r = p.add_run(text)
    r.bold           = bold
    r.italic         = italic
    r.font.name      = font
    r.font.size      = Pt(size_pt)
    r.font.color.rgb = color
    return r


def _no_borders_table(tbl):
    """Strip all visible borders from a python-docx Table."""
    tbl_el = tbl._tbl
    tpr = tbl_el.find(qn('w:tblPr'))
    if tpr is None:
        tpr = OxmlElement('w:tblPr')
        tbl_el.insert(0, tpr)
    bdr = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:val'), 'none')
        bdr.append(el)
    tpr.append(bdr)
    # Zero cell margins
    mar = OxmlElement('w:tblCellMar')
    for m in ('top', 'left', 'bottom', 'right'):
        mel = OxmlElement(f'w:{m}')
        mel.set(qn('w:w'), '0')
        mel.set(qn('w:type'), 'dxa')
        mar.append(mel)
    tpr.append(mar)


def _two_col_tbl(doc, left_txt, right_txt, font, usable_cm,
                 left_bold=False, left_italic=False,
                 left_size=10.5, right_size=9.5,
                 left_color=_BLACK_RGB, right_color=_GRAY_RGB,
                 space_before=20, space_after=0):
    """
    Insert a borderless 2-column table: bold title left, date right-aligned.
    More reliable than tab stops across different Word versions.
    """
    tbl = doc.add_table(rows=1, cols=2)
    tbl.style = 'Table Grid'
    _no_borders_table(tbl)

    lw = usable_cm * 0.72
    rw = usable_cm * 0.28
    row = tbl.rows[0]
    row.cells[0].width = Cm(lw)
    row.cells[1].width = Cm(rw)

    # Left cell
    pl = row.cells[0].paragraphs[0]
    pl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _d_spacing(pl, before=space_before, after=space_after)
    _d_run(pl, left_txt, font, left_size,
           bold=left_bold, italic=left_italic, color=left_color)

    # Right cell
    pr = row.cells[1].paragraphs[0]
    pr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _d_spacing(pr, before=space_before, after=space_after)
    _d_run(pr, right_txt, font, right_size, color=right_color)

    return tbl


def build_docx(resume_html: str,
               heading_font_css: str = 'Calibri',
               body_font_css:    str = 'Calibri') -> bytes:

    soup = BeautifulSoup(resume_html, 'html.parser')
    doc  = Document()
    BF   = _docx_font(body_font_css)
    HF   = _docx_font(heading_font_css)

    # A4 page
    section = doc.sections[0]
    section.page_width    = Cm(21.0)
    section.page_height   = Cm(29.7)
    section.left_margin   = Cm(1.8)
    section.right_margin  = Cm(1.8)
    section.top_margin    = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    USABLE = 17.4  # cm

    # ── Header ───────────────────────────────────────────────────────────────
    hdr = soup.find(class_='resume-header')
    if hdr:
        h1 = hdr.find('h1')
        if h1:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _d_spacing(p, before=0, after=10)
            _d_run(p, _get_text(h1).upper(), HF, 20,
                   bold=True, color=_NAVY_RGB)

        cbar = hdr.find(class_='contact-bar')
        if cbar:
            parts = [_get_text(a) for a in cbar.find_all('a') if _get_text(a)]
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _d_spacing(p, before=0, after=6)
            _d_run(p, '  \u2022  '.join(parts), BF, 9, color=_GRAY_RGB)

        # Divider — zero spacing so Word's Normal style doesn't add phantom gap
        p = doc.add_paragraph()
        _d_spacing(p, before=0, after=0)
        # suppress contextual spacing
        pPr = p._p.get_or_add_pPr()
        cs = OxmlElement('w:contextualSpacing')
        cs.set(qn('w:val'), '0')
        pPr.append(cs)
        _d_border_bottom(p)

    # ── Sections ─────────────────────────────────────────────────────────────
    for sec_el in soup.find_all(class_='section'):

        title_el = sec_el.find(class_='section-title')
        if title_el:
            p = doc.add_paragraph()
            _d_spacing(p, before=50, after=10)
            _d_run(p, _get_text(title_el).upper(), BF, 8.5,
                   bold=True, color=_NAVY_RGB)
            _d_border_bottom(p)

        sk = sec_el.find(class_='skills-text')
        if sk:
            p = doc.add_paragraph()
            _d_spacing(p, before=0, after=20)
            _d_run(p, _get_text(sk), BF, 10.5)

        for p_el in sec_el.find_all('p', recursive=False):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _d_spacing(p, before=0, after=20)
            _d_run(p, _get_text(p_el), BF, 10.5)

        # jobs
        for job in sec_el.find_all(class_='job'):
            jh = job.find(class_='job-header')
            if jh:
                _two_col_tbl(
                    doc,
                    _get_text(jh.find(class_='job-title')),
                    _get_text(jh.find(class_='job-date')),
                    BF, USABLE, left_bold=True, left_size=10.5,
                )
            cmp = job.find(class_='job-company')
            if cmp:
                p = doc.add_paragraph()
                _d_spacing(p, before=0, after=8)
                _d_run(p, _get_text(cmp), BF, 10,
                       italic=True, color=_IGRAY_RGB)
            for li in job.find_all('li'):
                p = doc.add_paragraph(style='List Bullet')
                _d_spacing(p, before=0, after=0)
                _d_run(p, _get_text(li), BF, 10.5)
            p = doc.add_paragraph()
            _d_spacing(p, before=0, after=8)

        # education
        for edu in sec_el.find_all(class_='edu-entry'):
            eh = edu.find(class_='edu-header')
            if eh:
                _two_col_tbl(
                    doc,
                    _get_text(eh.find(class_='edu-degree')),
                    _get_text(eh.find(class_='edu-date')),
                    BF, USABLE, left_bold=True, left_size=10.5,
                )
            es = edu.find(class_='edu-subheader')
            if es:
                _two_col_tbl(
                    doc,
                    _get_text(es.find(class_='edu-school')),
                    _get_text(es.find(class_='edu-gpa')),
                    BF, USABLE,
                    left_italic=True, left_size=10,
                    right_size=9.5, right_color=_NAVY_RGB,
                    space_before=0,
                )
            p = doc.add_paragraph()
            _d_spacing(p, before=0, after=8)

        # projects
        for proj in sec_el.find_all(class_='proj-entry'):
            ph = proj.find(class_='proj-header')
            if ph:
                ng    = ph.find(class_='proj-name-group') or ph
                pn    = _get_text(ng.find(class_='proj-name'))
                gh_el = ph.find(class_='proj-gh-link')
                pd    = _get_text(ph.find(class_='proj-date'))
                label = f"{pn}  |  {_get_text(gh_el)}" if gh_el else pn
                _two_col_tbl(
                    doc, label, pd, BF, USABLE,
                    left_bold=True, left_size=10.5,
                )
            desc_el = proj.find(class_='proj-desc')
            if desc_el:
                p = doc.add_paragraph()
                _d_spacing(p, before=0, after=8)
                _d_run(p, _get_text(desc_el), BF, 10.5)

        # certifications
        for cert in sec_el.find_all(class_='cert-entry'):
            p = doc.add_paragraph(style='List Bullet')
            _d_spacing(p, before=0, after=0)
            _d_run(p, _get_text(cert), BF, 10.5)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
